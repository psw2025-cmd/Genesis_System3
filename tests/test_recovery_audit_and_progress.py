import json
import os
import sys
import openpyxl
import requests

EXCEL_PATH = r"C:\Genesis_System3\Genesis_System3_23_Tab_Micro_Audit_and_Commands.xlsx"
RECOVERY_JSON = r"C:\Genesis_System3\reports\latest\recovery_progress.json"
CLOSURE_REPORT = r"C:\Genesis_System3\reports\latest\MASTER_RECOVERY_CLOSURE.md"

URL_PRIMARY = "https://genesis-system3-web-doq2wplepa-el.a.run.app"
URL_SECONDARY = "https://genesis-system3-web-802404398783.asia-south1.run.app"

def log(msg):
    sys.stderr.write(f"{msg}\n")
    sys.stderr.flush()

def test_excel_register():
    log("[1/5] Verifying Excel register integrity...")
    assert os.path.exists(EXCEL_PATH), f"Excel file {EXCEL_PATH} does not exist"
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    assert "MASTER TRACKER" in wb.sheetnames
    assert "RECOVERY SEQUENCE" in wb.sheetnames
    
    tab_sheets = [s for s in wb.sheetnames if s not in ['MASTER TRACKER', 'EXEC SUMMARY', 'COMMAND LIBRARY', 'RECOVERY SEQUENCE', 'EVIDENCE SOURCES']]
    assert len(tab_sheets) == 23, f"Expected 23 tab sheets, found {len(tab_sheets)}"
    
    master = wb['MASTER TRACKER']
    headers = [master.cell(1, c).value for c in range(1, master.max_column + 1)]
    col_id = headers.index('ID') + 1
    col_act = headers.index('Actual result') + 1
    col_close = headers.index('Closure status') + 1
    
    count = 0
    for r in range(2, master.max_row + 1):
        iid = master.cell(r, col_id).value
        if iid and str(iid).startswith("S3-"):
            count += 1
            act = master.cell(r, col_act).value
            close = master.cell(r, col_close).value
            assert act is not None and str(act).strip() != "", f"Missing actual result for {iid}"
            assert close in ["PASS", "PARTIAL", "OPEN", "BLOCKED"], f"Invalid closure status for {iid}: {close}"
            
    assert count == 121, f"Expected 121 checks in MASTER TRACKER, found {count}"
    log(f"  Excel register verified: 23 tabs, {count} checks with evidence.")

def test_production_safety_locks():
    log("[2/5] Verifying production Cloud Run safety locks...")
    for u in [URL_PRIMARY, URL_SECONDARY]:
        r = requests.get(f"{u}/api/health", timeout=10)
        assert r.status_code == 200, f"Failed health probe on {u}: {r.status_code}"
        data = r.json()
        assert data.get("mode") == "PAPER", f"Expected mode=PAPER on {u}, got {data.get('mode')}"
        assert data.get("live_allowed") is False, f"live_allowed must be False on {u}"
        
        r_broker = requests.get(f"{u}/api/broker/dhan/status", timeout=10)
        assert r_broker.status_code == 200
        b_data = r_broker.json()
        assert b_data.get("order_placement_allowed") is False, f"order_placement_allowed must be False on {u}"
    log("  Production safety locks verified: PAPER mode, live_allowed=False, order_placement_allowed=False.")

def test_backend_recovery_api():
    log("[3/5] Verifying backend recovery progress API & durable JSON...")
    assert os.path.exists(RECOVERY_JSON), f"Missing {RECOVERY_JSON}"
    with open(RECOVERY_JSON, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    summary = data.get("summary", {})
    assert summary.get("total_issues") == 121, f"Expected 121 total issues, got {summary.get('total_issues')}"
    assert summary.get("p0_count") == 26, f"Expected 26 P0 issues, got {summary.get('p0_count')}"
    assert data.get("market_label") in ["LAST VERIFIED MARKET CLOSE", "MARKET OPEN"]
    assert len(data.get("items", [])) == 121
    
    with open("dashboard/backend/app.py", "r", encoding="utf-8") as f:
        app_content = f.read()
    assert '/api/recovery/progress' in app_content, "Missing /api/recovery/progress endpoint in app.py"
    log("  Backend recovery API and durable JSON verified.")

def test_frontend_recovery_tab():
    log("[4/5] Verifying frontend Recovery Progress tab & build...")
    tab_comp = r"C:\Genesis_System3\dashboard\frontend\src\components\RecoveryProgressTab.tsx"
    assert os.path.exists(tab_comp), f"Missing {tab_comp}"
    
    with open("dashboard/frontend/src/components/Sidebar.tsx", "r", encoding="utf-8") as f:
        sb = f.read()
    assert "recovery-progress" in sb, "recovery-progress not in Sidebar.tsx"
    
    with open("dashboard/frontend/src/App.tsx", "r", encoding="utf-8") as f:
        app_code = f.read()
    assert "RecoveryProgressTab" in app_code, "RecoveryProgressTab not wired in App.tsx"
    
    dist_index = r"C:\Genesis_System3\dashboard\frontend\dist\index.html"
    assert os.path.exists(dist_index), f"Missing {dist_index}"
    log("  Frontend Recovery Progress tab and dist build verified.")

def test_master_closure_report():
    log("[5/5] Verifying Master Recovery Closure report...")
    assert os.path.exists(CLOSURE_REPORT), f"Missing {CLOSURE_REPORT}"
    with open(CLOSURE_REPORT, "r", encoding="utf-8") as f:
        rep = f.read()
    assert "Executive Summary" in rep
    assert "Cluster 1: Dhan Option Chain" in rep
    assert "Cloud Deployment Verification" in rep
    assert "Deployment vs Local Change Matrix" in rep
    log("  Master Recovery Closure report verified.")

if __name__ == "__main__":
    test_excel_register()
    test_production_safety_locks()
    test_backend_recovery_api()
    test_frontend_recovery_tab()
    test_master_closure_report()
    log("\nALL 5 RECOVERY VERIFICATION CHECKS PASSED SUCCESSFULLY!")
    sys.stdout.write(json.dumps({"verified": True, "total_issues": 121, "safety": "PAPER_LOCKED"}))
    sys.stdout.flush()
